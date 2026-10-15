"""Clinical batch screening report generator for PneumoVision.

Exports batch radiograph screening results to professionally formatted
PDF and DOCX (Word) documents with 3-column patient triage telemetry.
"""

from datetime import datetime
import io
from typing import List, Dict, Any


def generate_pdf_report(records: List[Dict[str, Any]], title: str = "PneumoVision Batch Screening Report") -> bytes:
    """Generate a clean, publication-grade PDF report of batch screening results."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        KeepTogether,
        HRFlowable,
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    # Custom typography styles
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0F172A"),
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#475569"),
        spaceAfter=12,
    )
    meta_style = ParagraphStyle(
        "MetaText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#334155"),
    )
    table_header_style = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#0F172A"),
        alignment=1,  # Center
    )
    patient_id_style = ParagraphStyle(
        "PatientIdText",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#0F172A"),
    )
    image_name_style = ParagraphStyle(
        "ImageNameText",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#64748B"),
    )
    pos_style = ParagraphStyle(
        "PosResult",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#DC2626"),
        alignment=1,
    )
    neg_style = ParagraphStyle(
        "NegResult",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#059669"),
        alignment=1,
    )
    density_style = ParagraphStyle(
        "DensityText",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#1E293B"),
        alignment=1,
    )
    disclaimer_style = ParagraphStyle(
        "Disclaimer",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#64748B"),
        alignment=1,
    )

    elements = []

    # Title & Metadata Header
    elements.append(Paragraph(title, title_style))
    elements.append(Paragraph("AI-Assisted Chest Radiograph Triage & Opacity Telemetry", subtitle_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#0284C7"), spaceAfter=10))

    # Summary Statistics Bar
    total_scans = len(records)
    pos_count = sum(1 for r in records if r.get("is_positive", False))
    neg_count = total_scans - pos_count
    pos_pct = (pos_count / total_scans * 100.0) if total_scans > 0 else 0.0

    timestamp_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    summary_text = (
        f"<b>Generated:</b> {timestamp_str} &nbsp;|&nbsp; "
        f"<b>Total Evaluated Scans:</b> {total_scans} &nbsp;|&nbsp; "
        f"<b>Positive (Suspected Pneumonia):</b> <font color='#DC2626'>{pos_count} ({pos_pct:.1f}%)</font> &nbsp;|&nbsp; "
        f"<b>Negative (Clear):</b> <font color='#059669'>{neg_count}</font>"
    )
    elements.append(Paragraph(summary_text, meta_style))
    elements.append(Spacer(1, 14))

    # Construct the 3-Column Table
    # Column 1: Patient ID (with image name below)
    # Column 2: Positive or Negative
    # Column 3: Density Index
    table_data = [
        [
            Paragraph("Patient ID & Radiograph", table_header_style),
            Paragraph("Screening Result", table_header_style),
            Paragraph("Pulmonary Density Index", table_header_style),
        ]
    ]

    for idx, r in enumerate(records):
        pid = str(r.get("patient_id", f"Patient-{idx+1}"))
        img_name = str(r.get("image_name", "N/A"))
        is_pos = r.get("is_positive", False)
        density = float(r.get("density_index", 0.0))

        # Cell 1: Patient ID with Image Name below it
        cell1 = [
            Paragraph(f"<b>{pid}</b>", patient_id_style),
            Spacer(1, 2),
            Paragraph(f"Image: {img_name}", image_name_style),
        ]

        # Cell 2: Positive or Negative result with clear indicator
        if is_pos:
            cell2 = Paragraph("POSITIVE<br/><font size='7.5' color='#991B1B'>Suspected Infiltrate</font>", pos_style)
        else:
            cell2 = Paragraph("NEGATIVE<br/><font size='7.5' color='#065F46'>Clear Lung Fields</font>", neg_style)

        # Cell 3: Density Index
        cutoff_note = "High Density" if is_pos else "Normal Density"
        cell3 = Paragraph(f"<b>{density:.3f}</b><br/><font size='7.5' color='#64748B'>Cutoff 0.300 ({cutoff_note})</font>", density_style)

        table_data.append([cell1, cell2, cell3])

    # Table Geometry & Layout
    # Total available width: 612 - 72 = 540 pt
    col_widths = [240, 150, 150]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)

    t_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("BOX", (0, 0), (-1, -1), 1.2, colors.HexColor("#CBD5E1")),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]

    # Alternating row shading
    for i in range(1, len(table_data)):
        if i % 2 == 0:
            t_style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F8FAFC")))

    t.setStyle(TableStyle(t_style))
    elements.append(t)
    elements.append(Spacer(1, 18))

    # Clinical Disclaimer & Guidelines
    elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#CBD5E1"), spaceAfter=8))
    elements.append(Paragraph(
        "<b>Medical Protocol Notice:</b> PneumoVision screening indices reflect DenseNet-121 operational feature activations "
        "trained on the RSNA Pneumonia Detection Challenge benchmark. Pulmonary Density Index ≥ 0.300 represents an operational "
        "sensitivity of 90.0% (F2=0.841). This automated report is designed for triage assistance and must be correlated with clinical symptoms "
        "and confirmed by a certified radiologist or licensed medical doctor.",
        disclaimer_style,
    ))

    doc.build(elements)
    return buffer.getvalue()


def generate_docx_report(records: List[Dict[str, Any]], title: str = "PneumoVision Batch Screening Report") -> bytes:
    """Generate a clean, formatted DOCX (Word) document of batch screening results."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn

    doc = Document()

    # Set standard page margins (0.6 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Document Title
    p_title = doc.add_paragraph()
    run_title = p_title.add_run(title)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)  # #0F172A
    p_title.paragraph_format.space_after = Pt(2)

    # Subtitle
    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run("AI-Assisted Chest Radiograph Triage & Opacity Telemetry")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)  # #475569
    p_sub.paragraph_format.space_after = Pt(10)

    # Summary Stats Bar
    total_scans = len(records)
    pos_count = sum(1 for r in records if r.get("is_positive", False))
    neg_count = total_scans - pos_count
    pos_pct = (pos_count / total_scans * 100.0) if total_scans > 0 else 0.0

    p_stats = doc.add_paragraph()
    p_stats.paragraph_format.space_after = Pt(14)
    run_stat1 = p_stats.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}   |   ")
    run_stat1.font.size = Pt(9.5)
    run_stat1.font.color.rgb = RGBColor(51, 65, 85)

    run_stat2 = p_stats.add_run(f"Total Scans: {total_scans}   |   ")
    run_stat2.font.bold = True
    run_stat2.font.size = Pt(9.5)

    run_pos = p_stats.add_run(f"Positive: {pos_count} ({pos_pct:.1f}%)   |   ")
    run_pos.font.bold = True
    run_pos.font.size = Pt(9.5)
    run_pos.font.color.rgb = RGBColor(220, 38, 38)  # Red

    run_neg = p_stats.add_run(f"Negative: {neg_count}")
    run_neg.font.bold = True
    run_neg.font.size = Pt(9.5)
    run_neg.font.color.rgb = RGBColor(5, 150, 105)  # Green

    # Construct the 3-Column Table
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(3.3), Inches(2.0), Inches(2.0)]
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    # Style Header Row
    hdr_cells = table.rows[0].cells
    hdr_titles = ["Patient ID & Radiograph", "Screening Result", "Pulmonary Density Index"]
    for i, title_text in enumerate(hdr_titles):
        cell = hdr_cells[i]
        cell.width = col_widths[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title_text)
        run.font.name = "Calibri"
        run.font.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(15, 23, 42)
        # Background color for header
        shading = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)

    # Populate Rows
    for idx, r in enumerate(records):
        row_cells = table.add_row().cells
        pid = str(r.get("patient_id", f"Patient-{idx+1}"))
        img_name = str(r.get("image_name", "N/A"))
        is_pos = r.get("is_positive", False)
        density = float(r.get("density_index", 0.0))

        # Cell 1: Patient ID & Image Name
        c1 = row_cells[0]
        c1.width = col_widths[0]
        c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        r_pid = p1.add_run(pid)
        r_pid.font.bold = True
        r_pid.font.size = Pt(10.5)
        r_pid.font.color.rgb = RGBColor(15, 23, 42)

        p1_sub = c1.add_paragraph()
        p1_sub.paragraph_format.space_after = Pt(0)
        r_img = p1_sub.add_run(f"Image: {img_name}")
        r_img.font.italic = True
        r_img.font.size = Pt(8.5)
        r_img.font.color.rgb = RGBColor(100, 116, 139)

        # Cell 2: Result
        c2 = row_cells[1]
        c2.width = col_widths[1]
        c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(1)
        if is_pos:
            r_res = p2.add_run("POSITIVE")
            r_res.font.bold = True
            r_res.font.size = Pt(11)
            r_res.font.color.rgb = RGBColor(220, 38, 38)
            p2_sub = c2.add_paragraph()
            p2_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_note = p2_sub.add_run("Suspected Infiltrate")
            r_note.font.size = Pt(8)
            r_note.font.color.rgb = RGBColor(153, 27, 27)
        else:
            r_res = p2.add_run("NEGATIVE")
            r_res.font.bold = True
            r_res.font.size = Pt(11)
            r_res.font.color.rgb = RGBColor(5, 150, 105)
            p2_sub = c2.add_paragraph()
            p2_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_note = p2_sub.add_run("Clear Lung Fields")
            r_note.font.size = Pt(8)
            r_note.font.color.rgb = RGBColor(6, 95, 70)

        # Cell 3: Density Index
        c3 = row_cells[2]
        c3.width = col_widths[2]
        c3.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p3 = c3.paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(1)
        r_den = p3.add_run(f"{density:.3f}")
        r_den.font.bold = True
        r_den.font.size = Pt(11)
        r_den.font.color.rgb = RGBColor(30, 41, 59)

        p3_sub = c3.add_paragraph()
        p3_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note_str = "Cutoff 0.300 (High)" if is_pos else "Cutoff 0.300 (Normal)"
        r_cut = p3_sub.add_run(note_str)
        r_cut.font.size = Pt(8)
        r_cut.font.color.rgb = RGBColor(100, 116, 139)

        # Alternating row background shading
        if idx % 2 == 1:
            for cell in row_cells:
                shd = parse_xml(r'<w:shd {} w:fill="F8FAFC"/>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(shd)

    # Disclaimer Note
    doc.add_paragraph()
    p_disc = doc.add_paragraph()
    p_disc.paragraph_format.space_before = Pt(14)
    r_disc = p_disc.add_run(
        "Medical Protocol Notice: PneumoVision screening indices reflect DenseNet-121 operational feature activations "
        "trained on the RSNA Pneumonia Detection Challenge benchmark. Pulmonary Density Index ≥ 0.300 represents an operational "
        "sensitivity of 90.0% (F2=0.841). This automated report is designed for triage assistance and must be correlated with clinical symptoms "
        "and confirmed by a certified radiologist or licensed medical doctor."
    )
    r_disc.font.size = Pt(8)
    r_disc.font.italic = True
    r_disc.font.color.rgb = RGBColor(100, 116, 139)

    out_buf = io.BytesIO()
    doc.save(out_buf)
    return out_buf.getvalue()
